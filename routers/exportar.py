from io import BytesIO

from fastapi import APIRouter, Cookie
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.auth import condominio_actual_id, no_permisos_response, puede_exportar, require_login
from core.database import get_conn


router = APIRouter()


@router.get("/exportar/visitas")
def exportar_visitas(admin_session: str | None = Cookie(default=None)):
    usuario = require_login(admin_session)
    if not puede_exportar(usuario):
        return no_permisos_response(usuario)
    with get_conn() as conn:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                SELECT v.id, v.nombre, v.rut, v.patente, d.torre, d.numero, v.autorizado_por,
                       v.observacion, v.hora_ingreso, v.hora_salida
                FROM visitas v
                LEFT JOIN departamentos d ON v.departamento_id = d.id
                WHERE v.condominio_id = %s
                ORDER BY v.id DESC
                """,
                (condominio_actual_id(usuario),),
            )
            visitas = cursor.fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Visitas"

    headers = [
        "ID",
        "Nombre visita",
        "RUT",
        "Patente",
        "Torre",
        "Departamento",
        "Autorizado por",
        "Observación",
        "Hora ingreso",
        "Hora salida",
    ]
    ws.append(headers)

    fill = PatternFill(fill_type="solid", fgColor="2563EB")
    font = Font(color="FFFFFF", bold=True)
    align = Alignment(horizontal="center")

    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = align

    for visita in visitas:
        ws.append(list(visita))

    for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J"]:
        ws.column_dimensions[col].width = 22

    archivo = BytesIO()
    wb.save(archivo)
    archivo.seek(0)

    return StreamingResponse(
        archivo,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=visitas_condominio.xlsx"},
    )


@router.get("/exportar/encomiendas")
def exportar_encomiendas(admin_session: str | None = Cookie(default=None)):
    usuario = require_login(admin_session)
    if not puede_exportar(usuario):
        return no_permisos_response(usuario)
    with get_conn() as conn:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                SELECT e.id, e.nombre_receptor, d.torre, d.numero, e.descripcion, e.recibido_por,
                       e.fecha_recepcion, e.entregado, e.fecha_entrega, e.entregado_a, e.observacion
                FROM encomiendas e
                LEFT JOIN departamentos d ON e.departamento_id = d.id
                WHERE e.condominio_id = %s
                ORDER BY e.id DESC
                """,
                (condominio_actual_id(usuario),),
            )
            data = cursor.fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Encomiendas"

    headers = [
        "ID",
        "Nombre receptor",
        "Torre",
        "Departamento",
        "Descripción",
        "Recibido por",
        "Fecha recepción",
        "Entregado",
        "Fecha entrega",
        "Entregado a",
        "Observación",
    ]
    ws.append(headers)

    fill = PatternFill(fill_type="solid", fgColor="2563EB")
    font = Font(color="FFFFFF", bold=True)
    align = Alignment(horizontal="center")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = align

    for row in data:
        ws.append(list(row))

    for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K"]:
        ws.column_dimensions[col].width = 22

    archivo = BytesIO()
    wb.save(archivo)
    archivo.seek(0)

    return StreamingResponse(
        archivo,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=encomiendas_condominio.xlsx"},
    )
