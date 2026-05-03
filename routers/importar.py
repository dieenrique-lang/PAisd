from io import BytesIO

from fastapi import APIRouter, Cookie, File, UploadFile
from fastapi.responses import HTMLResponse
from openpyxl import load_workbook

from core.auth import (
    condominio_actual_id,
    no_permisos_response,
    puede_escribir_residentes,
    puede_escribir_vehiculos,
    require_login,
)
from core.database import get_conn, obtener_o_crear_departamento
from core.helpers import encabezados_normalizados, h
from core.layout import render_resultado_importacion


router = APIRouter()


@router.post("/importar/residentes")
async def importar_residentes(admin_session: str | None = Cookie(default=None), archivo: UploadFile = File(...)):
    usuario = require_login(admin_session)
    if not puede_escribir_residentes(usuario):
        return no_permisos_response(usuario)

    if not archivo.filename or not archivo.filename.lower().endswith(".xlsx"):
        return HTMLResponse("Archivo inválido. Debe ser .xlsx", status_code=400)

    importados = 0
    omitidos = 0
    errores: list[str] = []

    try:
        contenido = await archivo.read()
        wb = load_workbook(filename=BytesIO(contenido), data_only=True)
        ws = wb.active
    except Exception as exc:
        return HTMLResponse(f"No se pudo leer el archivo: {h(exc)}", status_code=400)

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return HTMLResponse("El archivo está vacío.", status_code=400)

    headers = encabezados_normalizados(rows[0])
    required = ["nombre", "telefono", "email", "tipo", "torre", "numero"]
    if headers[: len(required)] != required and set(required) - set(headers):
        return HTMLResponse("Encabezados inválidos. Usa: nombre, telefono, email, tipo, torre, numero", status_code=400)

    idx = {hname: headers.index(hname) for hname in required}

    with get_conn() as conn:
        with conn.cursor() as cursor:
            for fila_num, row in enumerate(rows[1:], start=2):
                try:
                    row = row or ()
                    vals = [row[idx[k]] if idx[k] < len(row) else None for k in required]
                    nombre, telefono, email, tipo, torre, numero = [(str(v).strip() if v is not None else "") for v in vals]

                    if not any([nombre, telefono, email, tipo, torre, numero]):
                        omitidos += 1
                        continue
                    if not nombre or not numero:
                        omitidos += 1
                        errores.append(f"Fila {fila_num}: nombre y numero son obligatorios.")
                        continue
                    if not tipo:
                        tipo = "Residente"

                    dep_id = obtener_o_crear_departamento(cursor, condominio_actual_id(usuario), torre, numero)
                    cursor.execute(
                        """
                        INSERT INTO residentes (nombre, telefono, email, tipo, departamento_id, condominio_id)
                        VALUES (%s, %s, %s, %s, %s, %s)
                        """,
                        (nombre, telefono, email, tipo, dep_id, condominio_actual_id(usuario)),
                    )
                    conn.commit()
                    importados += 1
                except Exception as exc:
                    conn.rollback()
                    omitidos += 1
                    errores.append(f"Fila {fila_num}: {exc}")

    return render_resultado_importacion("Residentes", "/residentes", importados, omitidos, errores, usuario)


@router.post("/importar/vehiculos")
async def importar_vehiculos(admin_session: str | None = Cookie(default=None), archivo: UploadFile = File(...)):
    usuario = require_login(admin_session)
    if not puede_escribir_vehiculos(usuario):
        return no_permisos_response(usuario)

    if not archivo.filename or not archivo.filename.lower().endswith(".xlsx"):
        return HTMLResponse("Archivo inválido. Debe ser .xlsx", status_code=400)

    importados = 0
    omitidos = 0
    errores: list[str] = []

    try:
        contenido = await archivo.read()
        wb = load_workbook(filename=BytesIO(contenido), data_only=True)
        ws = wb.active
    except Exception as exc:
        return HTMLResponse(f"No se pudo leer el archivo: {h(exc)}", status_code=400)

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return HTMLResponse("El archivo está vacío.", status_code=400)

    headers = encabezados_normalizados(rows[0])
    required = ["patente", "marca", "modelo", "color", "torre", "numero"]
    optional = ["estacionamiento"]
    if set(required) - set(headers):
        return HTMLResponse("Encabezados inválidos. Usa: patente, marca, modelo, color, torre, numero, estacionamiento", status_code=400)

    idx = {hname: headers.index(hname) for hname in required if hname in headers}
    idx_opt = {hname: headers.index(hname) for hname in optional if hname in headers}

    with get_conn() as conn:
        with conn.cursor() as cursor:
            for fila_num, row in enumerate(rows[1:], start=2):
                try:
                    row = row or ()
                    vals = [row[idx[k]] if idx[k] < len(row) else None for k in required]
                    patente, marca, modelo, color, torre, numero = [(str(v).strip() if v is not None else "") for v in vals]
                    estacionamiento = ""
                    if "estacionamiento" in idx_opt:
                        val_est = row[idx_opt["estacionamiento"]] if idx_opt["estacionamiento"] < len(row) else None
                        estacionamiento = str(val_est).strip() if val_est is not None else ""

                    if not any([patente, marca, modelo, color, torre, numero]):
                        omitidos += 1
                        continue
                    if not patente or not numero:
                        omitidos += 1
                        errores.append(f"Fila {fila_num}: patente y numero son obligatorios.")
                        continue

                    dep_id = obtener_o_crear_departamento(cursor, condominio_actual_id(usuario), torre, numero)
                    cursor.execute(
                        """
                        INSERT INTO vehiculos (patente, marca, modelo, color, estacionamiento, departamento_id, condominio_id)
                        VALUES (%s, %s, %s, %s, %s, %s, %s)
                        """,
                        (patente.upper(), marca, modelo, color, estacionamiento, dep_id, condominio_actual_id(usuario)),
                    )
                    conn.commit()
                    importados += 1
                except Exception as exc:
                    conn.rollback()
                    omitidos += 1
                    errores.append(f"Fila {fila_num}: {exc}")

    return render_resultado_importacion("Vehículos", "/vehiculos", importados, omitidos, errores, usuario)
