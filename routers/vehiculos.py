from io import BytesIO

from fastapi import APIRouter, Cookie, File, Form, Query, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse
from openpyxl import load_workbook

from core.auth import (
    condominio_actual_id,
    no_permisos_response,
    puede_escribir_vehiculos,
    puede_ver_vehiculos,
    require_login,
)
from core.database import conectar, obtener_o_crear_departamento
from core.helpers import encabezados_normalizados, format_depto, h, render_delete_action
from core.layout import layout, render_resultado_importacion


router = APIRouter()


@router.get("/vehiculos", response_class=HTMLResponse)
def vehiculos(q: str = Query(default=""), admin_session: str | None = Cookie(default=None)):
    usuario = require_login(admin_session)
    if not puede_ver_vehiculos(usuario):
        return no_permisos_response(usuario)
    es_admin = puede_escribir_vehiculos(usuario)
    condominio_id = condominio_actual_id(usuario)
    with conectar() as conn:
        with conn.cursor() as cursor:
            where_parts = ["v.condominio_id = %s"]
            params = [condominio_id]
            if q:
                like = f"%{q}%"
                where_parts.append(
                    """
                    (
                        v.patente ILIKE %s OR
                        v.marca ILIKE %s OR
                        v.modelo ILIKE %s OR
                        v.color ILIKE %s OR
                        v.estacionamiento ILIKE %s OR
                        COALESCE(d.torre, '') ILIKE %s OR
                        d.numero ILIKE %s OR
                        (COALESCE(d.torre, '') || '-' || d.numero) ILIKE %s
                    )
                    """
                )
                params.extend([like, like, like, like, like, like, like, like])
            where_sql = f"WHERE {' AND '.join(where_parts)}" if where_parts else ""
            cursor.execute(
                """
                SELECT v.id, v.patente, v.marca, v.modelo, v.color, v.estacionamiento, d.torre, d.numero
                FROM vehiculos v
                LEFT JOIN departamentos d ON v.departamento_id = d.id
                """
                + where_sql
                + """
                ORDER BY v.id DESC
                LIMIT 200
                """,
                params,
            )
            data = cursor.fetchall()

    filas = "".join(
        f"""
        <tr>
            <td>{h(v[1])}</td><td>{h(v[2])}</td><td>{h(v[3])}</td><td>{h(v[4])}</td><td>{format_depto(v[6], v[7])}</td><td>{h(v[5])}</td>
            <td>{render_delete_action(es_admin, f'/eliminar-vehiculo/{v[0]}', '¿Eliminar vehículo?')}</td>
        </tr>
        """
        for v in data
    )

    form_html = """
    <div class="hero"><h1>Vehículos</h1><p>Registro de vehículos.</p></div>
    <div class="card">
        <h2>Agregar vehículo</h2>
        <form action="/guardar-vehiculo" method="post">
            <label>Patente<input name="patente" placeholder="Patente" required></label>
            <label>Marca<input name="marca" placeholder="Marca"></label>
            <label>Modelo<input name="modelo" placeholder="Modelo"></label>
            <label>Color<input name="color" placeholder="Color"></label>
            <label>Torre / Block<input name="torre" placeholder="Torre / Block"></label>
            <label>Departamento<input name="numero" placeholder="Departamento" required></label>
            <label>Estacionamiento<input name="estacionamiento" placeholder="N° estacionamiento"></label>
            <button class="full" type="submit">Guardar vehículo</button>
        </form>
    </div>
    <div class="card">
        <h2>Importar vehículos desde Excel</h2>
        <p class="muted">Columnas requeridas: patente, marca, modelo, color, torre, numero, estacionamiento (estacionamiento es opcional)</p>
        <form action="/importar/vehiculos" method="post" enctype="multipart/form-data">
            <label>Archivo .xlsx<input type="file" name="archivo" accept=".xlsx" required></label>
            <button class="full" type="submit">Importar vehículos</button>
        </form>
    </div>
    """ if es_admin else """
    <div class="hero"><h1>Vehículos</h1><p>Vista de solo lectura para comité.</p></div>
    """

    contenido = f"""
    {form_html}
    <div class="card">
        <h2>Buscar y filtrar</h2>
        <form action="/vehiculos" method="get">
            <label>Búsqueda
                <input name="q" value="{h(q)}" placeholder="Buscar por patente, marca, modelo, color, estacionamiento o depto">
            </label>
            <button type="submit">Aplicar filtros</button>
            <a class="btn dark" href="/vehiculos">Limpiar</a>
        </form>
    </div>
    <div class="card">
        <h2>Listado vehículos</h2>
        <div class="table-wrap"><table>
            <tr><th>Patente</th><th>Marca</th><th>Modelo</th><th>Color</th><th>Depto</th><th>Estacionamiento</th><th>Acción</th></tr>
            {filas}
        </table></div>
    </div>
    <div class="actions"><a class="btn" href="/">Inicio</a></div>
    """
    return layout("Vehículos", contenido, usuario)


@router.post("/guardar-vehiculo")
def guardar_vehiculo(
    admin_session: str | None = Cookie(default=None),
    patente: str = Form(...),
    marca: str = Form(""),
    modelo: str = Form(""),
    color: str = Form(""),
    torre: str = Form(""),
    numero: str = Form(...),
    estacionamiento: str = Form(""),
):
    usuario = require_login(admin_session)
    if not puede_escribir_vehiculos(usuario):
        return no_permisos_response(usuario)
    with conectar() as conn:
        with conn.cursor() as cursor:
            dep_id = obtener_o_crear_departamento(cursor, condominio_actual_id(usuario), torre, numero)
            cursor.execute(
                """
                INSERT INTO vehiculos (patente, marca, modelo, color, estacionamiento, departamento_id, condominio_id)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                """,
                (patente.upper(), marca, modelo, color, estacionamiento, dep_id, condominio_actual_id(usuario)),
            )
        conn.commit()
    return RedirectResponse(url="/vehiculos", status_code=303)


@router.post("/importar/vehiculos")
async def importar_vehiculos(admin_session: str | None = Cookie(default=None), archivo: UploadFile = File(...)):
    usuario = require_login(admin_session)
    if not puede_escribir_vehiculos(usuario):
        return no_permisos_response(usuario)

    if not archivo.filename or not archivo.filename.lower().endswith(".xlsx"):
        return HTMLResponse("Archivo inválido. Debe ser .xlsx", status_code=400)

    importados = 0
    omitidos = 0
    errores: list[str] = []

    try:
        contenido = await archivo.read()
        wb = load_workbook(filename=BytesIO(contenido), data_only=True)
        ws = wb.active
    except Exception as exc:
        return HTMLResponse(f"No se pudo leer el archivo: {h(exc)}", status_code=400)

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return HTMLResponse("El archivo está vacío.", status_code=400)

    headers = encabezados_normalizados(rows[0])
    required = ["patente", "marca", "modelo", "color", "torre", "numero"]
    optional = ["estacionamiento"]
    if set(required) - set(headers):
        return HTMLResponse("Encabezados inválidos. Usa: patente, marca, modelo, color, torre, numero, estacionamiento", status_code=400)

    idx = {hname: headers.index(hname) for hname in required if hname in headers}
    idx_opt = {hname: headers.index(hname) for hname in optional if hname in headers}

    with conectar() as conn:
        with conn.cursor() as cursor:
            for fila_num, row in enumerate(rows[1:], start=2):
                try:
                    row = row or ()
                    vals = [row[idx[k]] if idx[k] < len(row) else None for k in required]
                    patente, marca, modelo, color, torre, numero = [(str(v).strip() if v is not None else "") for v in vals]
                    estacionamiento = ""
                    if "estacionamiento" in idx_opt:
                        val_est = row[idx_opt["estacionamiento"]] if idx_opt["estacionamiento"] < len(row) else None
                        estacionamiento = str(val_est).strip() if val_est is not None else ""

                    if not any([patente, marca, modelo, color, torre, numero]):
                        omitidos += 1
                        continue
                    if not patente or not numero:
                        omitidos += 1
                        errores.append(f"Fila {fila_num}: patente y numero son obligatorios.")
                        continue

                    dep_id = obtener_o_crear_departamento(cursor, condominio_actual_id(usuario), torre, numero)
                    cursor.execute(
                        """
                        INSERT INTO vehiculos (patente, marca, modelo, color, estacionamiento, departamento_id, condominio_id)
                        VALUES (%s, %s, %s, %s, %s, %s, %s)
                        """,
                        (patente.upper(), marca, modelo, color, estacionamiento, dep_id, condominio_actual_id(usuario)),
                    )
                    conn.commit()
                    importados += 1
                except Exception as exc:
                    conn.rollback()
                    omitidos += 1
                    errores.append(f"Fila {fila_num}: {exc}")

    return render_resultado_importacion("Vehículos", "/vehiculos", importados, omitidos, errores, usuario)


@router.get("/eliminar-vehiculo/{vehiculo_id}")
def eliminar_vehiculo(vehiculo_id: int, admin_session: str | None = Cookie(default=None)):
    usuario = require_login(admin_session)
    if not puede_escribir_vehiculos(usuario):
        return no_permisos_response(usuario)

    with conectar() as conn:
        with conn.cursor() as cursor:
            cursor.execute(
                "DELETE FROM vehiculos WHERE id = %s AND condominio_id = %s",
                (vehiculo_id, condominio_actual_id(usuario)),
            )
        conn.commit()
    return RedirectResponse(url="/vehiculos", status_code=303)
